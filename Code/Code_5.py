# ============================================================
# 0. IMPORT LIBRARIES
# ============================================================
import os
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import cartopy.crs as ccrs
import cartopy.feature as cfeature
from matplotlib.lines import Line2D
from openpyxl import load_workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)
# ============================================================
# 1. FILE SETTINGS
# ============================================================
DATA_FILE = "Siklon_Tropis_ERA5_Extracted.xlsx"
CLUSTER_FILE = os.path.join(
    "output_3_1_2",
    "Clustering_Results_3_1_2_2.xlsx"
)
OUTPUT_DIR = "output_3_1_3"
os.makedirs(
    OUTPUT_DIR,
    exist_ok=True
)
OUTPUT_EXCEL = os.path.join(
    OUTPUT_DIR,
    "Equatorial_Rare_Event_Analysis_3_1_3_4.xlsx"
)
OUTPUT_G314 = os.path.join(
    OUTPUT_DIR,
    "Figure_3.14_Equatorial_Rare_Event_Trajectories.png"
)
OUTPUT_G315 = os.path.join(
    OUTPUT_DIR,
    "Figure_3.15_Equatorial_Rare_Event_Meteorology.png"
)
# ============================================================
# 2. STUDY DOMAIN
# ============================================================
LON_MIN = 95
LON_MAX = 145
LAT_MIN = -15
LAT_MAX = 25
EQ_MIN = -5
EQ_MAX = 5
YEAR_START = 2001
YEAR_END = 2025
# ============================================================
# 3. CASE COLORS
# ============================================================
CASE_COLORS = {
    "VAMEI":
        "#1f77b4",
    "SENYAR":
        "#ff7f0e"
}
# ============================================================
# 4. VISUAL SETTINGS
# ============================================================
plt.rcParams.update({
    "font.family":
        "DejaVu Sans",
    "font.size":
        10,
    "axes.titlesize":
        11,
    "axes.labelsize":
        10,
    "figure.dpi":
        150,
    "savefig.dpi":
        300
})
# ============================================================
# 5. CHECK INPUT FILE
# ============================================================
if not os.path.exists(DATA_FILE):
    raise FileNotFoundError(
        f"File not found:\n{DATA_FILE}"
    )
# ============================================================
# 6. READ INPUT DATA
# ============================================================
df = pd.read_excel(
    DATA_FILE,
    sheet_name="Sheet1",
    dtype={
        "SID": str
    }
)
df["SID"] = (
    df["SID"]
    .astype(str)
    .str.strip()
)
df["ISO_TIME"] = pd.to_datetime(
    df["ISO_TIME"],
    errors="coerce"
)
numeric_columns = [
    "LAT",
    "LON",
    "SEASON",
    "WIND_km/h",
    "PRESSURE_hPa",
    "era5_sst",
    "era5_tp",
    "era5_w",
    "era5_vo"
]
for col in numeric_columns:
    df[col] = pd.to_numeric(
        df[col],
        errors="coerce"
    )
# ============================================================
# 7. CLEANING
# ============================================================
df = df.dropna(
    subset=[
        "SID",
        "ISO_TIME",
        "LAT",
        "LON"
    ]
).copy()
df = df[
    df["LAT"].between(
        -90,
        90
    )
].copy()
df = df.drop_duplicates(
    subset=[
        "SID",
        "ISO_TIME"
    ],
    keep="first"
)
df = df.sort_values(
    [
        "SID",
        "ISO_TIME"
    ]
).reset_index(
    drop=True
)
# ============================================================
# 8. IDENTIFY FIRST OBSERVATION FOR EVERY SID
# ============================================================
genesis_all = (
    df.drop_duplicates(
        subset="SID",
        keep="first"
    )
    .copy()
)
genesis_all = genesis_all[
    genesis_all["SEASON"].between(
        YEAR_START,
        YEAR_END
    )
].copy()
# ============================================================
# 9. FILTER GENESIS BY STUDY DOMAIN
# ============================================================
selected_genesis = genesis_all[
    genesis_all["LON"].between(
        LON_MIN,
        LON_MAX
    )
    &
    genesis_all["LAT"].between(
        LAT_MIN,
        LAT_MAX
    )
].copy()
# ============================================================
# 10. IDENTIFY EQUATORIAL RARE EVENTS
# ============================================================
equatorial_genesis = selected_genesis[
    selected_genesis["LAT"].between(
        EQ_MIN,
        EQ_MAX,
        inclusive="both"
    )
].copy()
print("=" * 80)
print(
    "RARE TROPICAL CYCLONE GENESIS EVENTS IN THE EQUATORIAL ZONE"
)
print("=" * 80)
print(
    equatorial_genesis[
        [
            "SID",
            "NAME",
            "SEASON",
            "ISO_TIME",
            "LAT",
            "LON"
        ]
    ]
    .to_string(
        index=False
    )
)
print(
    "\nNumber of rare events =",
    len(
        equatorial_genesis
    )
)
if len(
    equatorial_genesis
) != 2:
    print(
        "\nWARNING:"
        "\nThe number of equatorial cases is not 2."
    )
if equatorial_genesis.empty:
    raise ValueError(
        "No equatorial genesis cases were found within the selected study domain and period."
    )
# ============================================================
# 11. EXTRACT COMPLETE TRAJECTORIES OF EQUATORIAL CASES
# ============================================================
#
# Once genesis is within the study domain,
# the ENTIRE subsequent trajectory is retained.
# ============================================================
rare_sids = (
    equatorial_genesis[
        "SID"
    ]
    .tolist()
)
rare_tracks = df[
    df["SID"].isin(
        rare_sids
    )
].copy()
rare_tracks = rare_tracks.sort_values(
    [
        "SID",
        "ISO_TIME"
    ]
).reset_index(
    drop=True
)
# ============================================================
# 12. CONVERT METEOROLOGICAL PARAMETERS
# ============================================================
# SST
rare_tracks["SST_C"] = (
    rare_tracks["era5_sst"]
    -
    273.15
)
# Total precipitation
rare_tracks["TP_mm"] = (
    rare_tracks["era5_tp"]
    *
    1000.0
)
# Vertical velocity
rare_tracks["Omega_Pa_s"] = (
    rare_tracks["era5_w"]
)
# Relative vorticity
rare_tracks["Vorticity_1e5"] = (
    rare_tracks["era5_vo"]
    *
    1e5
)
# Wind Zoom Earth
rare_tracks["Wind_kmh"] = (
    rare_tracks["WIND_km/h"]
)
# Pressure Zoom Earth
rare_tracks["Pressure_hPa"] = (
    rare_tracks["PRESSURE_hPa"]
)
# Coriolis
OMEGA_EARTH = 7.2921159e-5
rare_tracks["Coriolis_1e5"] = (
    2
    *
    OMEGA_EARTH
    *
    np.sin(
        np.radians(
            rare_tracks["LAT"]
        )
    )
    *
    1e5
)
# ============================================================
# 13. CALCULATE RELATIVE TIME
# ============================================================
processed_tracks = []
for sid, g in rare_tracks.groupby(
    "SID"
):
    g = (
        g.sort_values(
            "ISO_TIME"
        )
        .copy()
    )
    start_time = (
        g["ISO_TIME"].iloc[0]
    )
    end_time = (
        g["ISO_TIME"].iloc[-1]
    )
    duration_hours = (
        (
            end_time
            -
            start_time
        )
        .total_seconds()
        /
        3600.0
    )
    g["Elapsed_Hours"] = (
        (
            g["ISO_TIME"]
            -
            start_time
        )
        .dt.total_seconds()
        /
        3600.0
    )
    if duration_hours > 0:
        g["Life_Cycle_Percent"] = (
            g["Elapsed_Hours"]
            /
            duration_hours
            *
            100
        )
    else:
        g["Life_Cycle_Percent"] = 0
    processed_tracks.append(
        g
    )
rare_tracks = pd.concat(
    processed_tracks,
    ignore_index=True
)
# ============================================================
# 14. HAVERSINE DISTANCE FUNCTION
# ============================================================
def haversine_km(
    lat1,
    lon1,
    lat2,
    lon2
):
    R = 6371.0088
    lat1 = np.radians(
        lat1
    )
    lon1 = np.radians(
        lon1
    )
    lat2 = np.radians(
        lat2
    )
    lon2 = np.radians(
        lon2
    )
    dlat = (
        lat2
        -
        lat1
    )
    dlon = (
        lon2
        -
        lon1
    )
    a = (
        np.sin(
            dlat / 2
        ) ** 2
        +
        np.cos(
            lat1
        )
        *
        np.cos(
            lat2
        )
        *
        np.sin(
            dlon / 2
        ) ** 2
    )
    c = (
        2
        *
        np.arcsin(
            np.sqrt(
                a
            )
        )
    )
    return (
        R
        *
        c
    )
# ============================================================
# 15. CALCULATE TRAJECTORY LENGTH
# ============================================================
def calculate_track_length(
    g
):
    g = (
        g.sort_values(
            "ISO_TIME"
        )
    )
    lat = (
        g["LAT"]
        .to_numpy()
    )
    lon = (
        g["LON"]
        .to_numpy()
    )
    if len(
        g
    ) < 2:
        return 0
    distances = haversine_km(
        lat[:-1],
        lon[:-1],
        lat[1:],
        lon[1:]
    )
    return float(
        np.nansum(
            distances
        )
    )
# ============================================================
# 16. READ CLUSTER ASSIGNMENTS
# ============================================================
if os.path.exists(
    CLUSTER_FILE
):
    cluster_assignment = pd.read_excel(
        CLUSTER_FILE,
        sheet_name="Cluster Assignments",
        dtype={
            "SID": str
        }
    )
    cluster_assignment["SID"] = (
        cluster_assignment["SID"]
        .astype(str)
        .str.strip()
    )
    cluster_assignment = (
        cluster_assignment[
            [
                "SID",
                "Cluster"
            ]
        ]
    )
else:
    cluster_assignment = pd.DataFrame(
        columns=[
            "SID",
            "Cluster"
        ]
    )
# ============================================================
# 17. TABLE 3.9
# METEOROLOGICAL CONDITIONS AT GENESIS
# ============================================================
genesis_rare = (
    rare_tracks
    .sort_values(
        [
            "SID",
            "ISO_TIME"
        ]
    )
    .drop_duplicates(
        subset="SID",
        keep="first"
    )
    .copy()
)
genesis_rare = genesis_rare.merge(
    cluster_assignment,
    on="SID",
    how="left"
)
table_39 = genesis_rare[
    [
        "SID",
        "NAME",
        "SEASON",
        "LAT",
        "LON",
        "Cluster",
        "SST_C",
        "TP_mm",
        "Omega_Pa_s",
        "Vorticity_1e5",
        "Wind_kmh",
        "Pressure_hPa",
        "Coriolis_1e5"
    ]
].copy()
table_39 = table_39.rename(
    columns={
        "NAME":
            "Cyclone",
        "SEASON":
            "Year",
        "LAT":
            "Genesis Latitude",
        "LON":
            "Genesis Longitude",
        "SST_C":
            "SST (°C)",
        "TP_mm":
            "TP (mm)",
        "Omega_Pa_s":
            "ω (Pa/s)",
        "Vorticity_1e5":
            "Relative Vorticity (×10⁻⁵ s⁻¹)",
        "Wind_kmh":
            "Wind (km/h)",
        "Pressure_hPa":
            "Pressure (hPa)",
        "Coriolis_1e5":
            "Coriolis Parameter (×10⁻⁵ s⁻¹)"
    }
)
# ============================================================
# 18. TABLE 3.10
# SUMMARY OF DEVELOPMENT ALONG THE TRAJECTORY
# ============================================================
summary_records = []
for sid, g in rare_tracks.groupby(
    "SID"
):
    g = (
        g.sort_values(
            "ISO_TIME"
        )
        .copy()
    )
    name = (
        g["NAME"].iloc[0]
    )
    year = int(
        g["SEASON"].iloc[0]
    )
    duration_hours = (
        g["Elapsed_Hours"]
        .max()
    )
    track_length = (
        calculate_track_length(
            g
        )
    )
    summary_records.append({
        "Cyclone":
            name,
        "Year":
            year,
        "Duration (hours)":
            duration_hours,
        "Track Length (km)":
            track_length,
        "Mean SST (°C)":
            g[
                "SST_C"
            ].mean(),
        "Mean Total Precipitation (mm)":
            g[
                "TP_mm"
            ].mean(),
        "Mean Vertical Velocity (Pa/s)":
            g[
                "Omega_Pa_s"
            ].mean(),
        "Minimum Vertical Velocity (Pa/s)":
            g[
                "Omega_Pa_s"
            ].min(),
        "Mean Relative Vorticity (×10⁻⁵ s⁻¹)":
            g[
                "Vorticity_1e5"
            ].mean(),
        "Maximum Relative Vorticity (×10⁻⁵ s⁻¹)":
            g[
                "Vorticity_1e5"
            ].max(),
        "Maximum Wind Speed (km/h)":
            g[
                "Wind_kmh"
            ].max(),
        "Minimum Pressure (hPa)":
            g[
                "Pressure_hPa"
            ].min()
    })
table_310 = pd.DataFrame(
    summary_records
)
# ============================================================
# 19. TIME OF MAXIMUM WIND AND MINIMUM PRESSURE
# ============================================================
intensity_records = []
for sid, g in rare_tracks.groupby(
    "SID"
):
    g = (
        g.sort_values(
            "ISO_TIME"
        )
        .copy()
    )
    name = (
        g["NAME"].iloc[0]
    )
    # Maximum wind speed (ignore missing observations)
    wind_valid = g[g["Wind_kmh"].notna()]
    wind_idx = wind_valid["Wind_kmh"].idxmax() if not wind_valid.empty else None

    # Minimum pressure (ignore missing observations)
    pressure_valid = g[g["Pressure_hPa"].notna()]
    pressure_idx = pressure_valid["Pressure_hPa"].idxmin() if not pressure_valid.empty else None

    intensity_records.append({
        "Cyclone": name,
        "Maximum Wind Speed": g.loc[wind_idx, "Wind_kmh"] if wind_idx is not None else np.nan,
        "Time of Maximum Wind Speed": g.loc[wind_idx, "ISO_TIME"] if wind_idx is not None else pd.NaT,
        "Life-Cycle Phase at Maximum Wind (%)": g.loc[wind_idx, "Life_Cycle_Percent"] if wind_idx is not None else np.nan,
        "Minimum Pressure": g.loc[pressure_idx, "Pressure_hPa"] if pressure_idx is not None else np.nan,
        "Time of Minimum Pressure": g.loc[pressure_idx, "ISO_TIME"] if pressure_idx is not None else pd.NaT,
        "Life-Cycle Phase at Minimum Pressure (%)": g.loc[pressure_idx, "Life_Cycle_Percent"] if pressure_idx is not None else np.nan
    })
intensity_df = pd.DataFrame(
    intensity_records
)
# ============================================================
# 20. PRINT TABLES
# ============================================================
print("\n" + "=" * 120)
print(
    "TABLE 3.9 - METEOROLOGICAL CONDITIONS AT GENESIS"
)
print("=" * 120)
print(
    table_39
    .round(3)
    .to_string(
        index=False
    )
)
print("\n" + "=" * 120)
print(
    "TABLE 3.10 - DEVELOPMENT SUMMARY"
)
print("=" * 120)
print(
    table_310
    .round(3)
    .to_string(
        index=False
    )
)
# ============================================================
# 21. SAVE RESULTS TO EXCEL
# ============================================================
with pd.ExcelWriter(
    OUTPUT_EXCEL,
    engine="openpyxl"
) as writer:
    table_39.round(
        4
    ).to_excel(
        writer,
        sheet_name="Table 3.9",
        index=False
    )
    table_310.round(
        4
    ).to_excel(
        writer,
        sheet_name="Table 3.10",
        index=False
    )
    intensity_df.round(
        4
    ).to_excel(
        writer,
        sheet_name="Intensity Phases",
        index=False
    )
    rare_tracks.to_excel(
        writer,
        sheet_name="Complete Time Series",
        index=False
    )
# ============================================================
# 22. FORMAT EXCEL WORKBOOK
# ============================================================
wb = load_workbook(
    OUTPUT_EXCEL
)
header_fill = PatternFill(
    "solid",
    fgColor="D9EAF7"
)
header_font = Font(
    bold=True
)
thin_border = Border(
    left=Side(
        style="thin",
        color="BFBFBF"
    ),
    right=Side(
        style="thin",
        color="BFBFBF"
    ),
    top=Side(
        style="thin",
        color="BFBFBF"
    ),
    bottom=Side(
        style="thin",
        color="BFBFBF"
    )
)
for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.fill = (
            header_fill
        )
        cell.font = (
            header_font
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
    for row in ws.iter_rows():
        for cell in row:
            cell.border = (
                thin_border
            )
            cell.alignment = Alignment(
                vertical="center"
            )
    for column_cells in ws.columns:
        column_letter = (
            column_cells[0]
            .column_letter
        )
        max_length = 0
        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(
                        str(
                            cell.value
                        )
                    )
                )
        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            40
        )
wb.save(
    OUTPUT_EXCEL
)
# ============================================================
# ============================================================
#
#              GAMBAR 3.14
#
# TRAJECTORIES OF RARE EQUATORIAL EVENTS
#
# ============================================================
# ============================================================
fig = plt.figure(
    figsize=(
        13,
        8
    )
)
ax = plt.axes(
    projection=ccrs.PlateCarree()
)
# ------------------------------------------------------------
# Regional map extent
# ------------------------------------------------------------
ax.set_extent(
    [
        88,
        110,
        -6,
        8
    ],
    crs=ccrs.PlateCarree()
)
# ------------------------------------------------------------
# Background
# ------------------------------------------------------------
ax.add_feature(
    cfeature.OCEAN,
    facecolor="#dfe8ef"
)
ax.add_feature(
    cfeature.LAND,
    facecolor="#f2f2f2"
)
ax.add_feature(
    cfeature.COASTLINE,
    linewidth=0.7
)
ax.add_feature(
    cfeature.BORDERS,
    linewidth=0.45,
    linestyle=":"
)
# ------------------------------------------------------------
# Equatorial zone boundaries
# ------------------------------------------------------------
for lat_line in [
    -5,
    5
]:
    ax.plot(
        [
            88,
            110
        ],
        [
            lat_line,
            lat_line
        ],
        transform=ccrs.PlateCarree(),
        color="gray",
        linestyle="--",
        linewidth=1.0
    )
# Equator
ax.plot(
    [
        88,
        110
    ],
    [
        0,
        0
    ],
    transform=ccrs.PlateCarree(),
    color="black",
    linestyle="-.",
    linewidth=1.0
)
# Western boundary of the study domain
ax.plot(
    [
        95,
        95
    ],
    [
        -6,
        8
    ],
    transform=ccrs.PlateCarree(),
    color="gray",
    linestyle=":",
    linewidth=1.2
)
# ------------------------------------------------------------
# Plot both cases
# ------------------------------------------------------------
for sid, g in rare_tracks.groupby(
    "SID"
):
    g = (
        g.sort_values(
            "ISO_TIME"
        )
    )
    name = str(
        g[
            "NAME"
        ].iloc[0]
    )
    year = int(
        g[
            "SEASON"
        ].iloc[0]
    )
    color = CASE_COLORS.get(
        name,
        "black"
    )
    # Track
    ax.plot(
        g["LON"],
        g["LAT"],
        transform=ccrs.PlateCarree(),
        color=color,
        linewidth=2.4,
        alpha=0.9,
        label=(
            f"{name} ({year})"
        )
    )
    # Track points
    ax.scatter(
        g["LON"],
        g["LAT"],
        transform=ccrs.PlateCarree(),
        color=color,
        s=12,
        alpha=0.35
    )
    # Genesis
    ax.scatter(
        g[
            "LON"
        ].iloc[0],
        g[
            "LAT"
        ].iloc[0],
        transform=ccrs.PlateCarree(),
        color=color,
        marker="o",
        s=90,
        edgecolor="black",
        linewidth=1.0,
        zorder=10
    )
    # Termination
    ax.scatter(
        g[
            "LON"
        ].iloc[-1],
        g[
            "LAT"
        ].iloc[-1],
        transform=ccrs.PlateCarree(),
        color=color,
        marker="X",
        s=100,
        edgecolor="black",
        linewidth=0.8,
        zorder=10
    )
    # Label cyclone name at genesis
    ax.text(
        g[
            "LON"
        ].iloc[0]
        +
        0.25,
        g[
            "LAT"
        ].iloc[0]
        +
        0.25,
        f"{name} ({year})",
        transform=ccrs.PlateCarree(),
        fontsize=10,
        fontweight="bold",
        color=color,
        bbox=dict(
            facecolor="white",
            edgecolor=color,
            alpha=0.85,
            boxstyle="round,pad=0.25"
        )
    )
# ------------------------------------------------------------
# Grid
# ------------------------------------------------------------
gl = ax.gridlines(
    draw_labels=True,
    linewidth=0.4,
    color="gray",
    alpha=0.5,
    linestyle="--"
)
gl.top_labels = False
gl.right_labels = False
# ------------------------------------------------------------
# Legend
# ------------------------------------------------------------
legend_cases = [
    Line2D(
        [0],
        [0],
        color=CASE_COLORS[
            "VAMEI"
        ],
        linewidth=2.5,
        label="VAMEI (2001)"
    ),
    Line2D(
        [0],
        [0],
        color=CASE_COLORS[
            "SENYAR"
        ],
        linewidth=2.5,
        label="SENYAR (2025)"
    ),
    Line2D(
        [0],
        [0],
        marker="o",
        color="white",
        markerfacecolor="gray",
        markeredgecolor="black",
        markersize=8,
        label="Genesis"
    ),
    Line2D(
        [0],
        [0],
        marker="X",
        color="white",
        markerfacecolor="gray",
        markeredgecolor="black",
        markersize=8,
        label="Termination"
    )
]
ax.legend(
    handles=legend_cases,
    loc="upper center",
    bbox_to_anchor=(0.5, -0.09),
    ncol=4,
    frameon=True,
    framealpha=0.95,
    fontsize=9,
    columnspacing=1.6,
    handlelength=2.4
)
# ------------------------------------------------------------
# Title
# ------------------------------------------------------------
ax.set_title(
    "Trajectories of Rare Tropical Cyclones with Equatorial Genesis",
    fontsize=16,
    fontweight="bold",
    pad=14
)
plt.tight_layout()
plt.subplots_adjust(bottom=0.16)
plt.savefig(
    OUTPUT_G314,
    dpi=300,
    bbox_inches="tight",
    facecolor="white"
)
plt.show()
# ============================================================
# ============================================================
#
#              GAMBAR 3.15
#
# EVOLUTION OF METEOROLOGICAL PARAMETERS
#
# ============================================================
# ============================================================
PLOT_PARAMETERS = [
    (
        "SST_C",
        "Sea Surface Temperature",
        "°C"
    ),
    (
        "TP_mm",
        "Total Precipitation",
        "mm"
    ),
    (
        "Omega_Pa_s",
        "Vertical Velocity (ω)",
        "Pa/s"
    ),
    (
        "Vorticity_1e5",
        "Relative Vorticity",
        "×10⁻⁵ s⁻¹"
    ),
    (
        "Wind_kmh",
        "Wind",
        "km/h"
    ),
    (
        "Pressure_hPa",
        "Pressure",
        "hPa"
    ),
    (
        "Coriolis_1e5",
        "Coriolis Parameter",
        "×10⁻⁵ s⁻¹"
    )
]
fig, axes = plt.subplots(
    2,
    4,
    figsize=(
        18,
        10
    )
)
axes = axes.flatten()
# ------------------------------------------------------------
# Plot
# ------------------------------------------------------------
for ax, (
    variable,
    title,
    unit
) in zip(
    axes,
    PLOT_PARAMETERS
):
    for sid, g in rare_tracks.groupby(
        "SID"
    ):
        g = (
            g.sort_values(
                "Life_Cycle_Percent"
            )
        )
        name = str(
            g[
                "NAME"
            ].iloc[0]
        )
        year = int(
            g[
                "SEASON"
            ].iloc[0]
        )
        color = CASE_COLORS.get(
            name,
            "black"
        )
        ax.plot(
            g[
                "Life_Cycle_Percent"
            ],
            g[
                variable
            ],
            color=color,
            linewidth=2.2,
            marker="o",
            markersize=2.5,
            alpha=0.9,
            label=(
                f"{name} ({year})"
            )
        )
    ax.set_title(
        title,
        fontweight="bold"
    )
    ax.set_xlabel(
        "Relative Life Cycle (%)"
    )
    ax.set_ylabel(
        unit
    )
    ax.set_xlim(
        0,
        100
    )
    ax.set_xticks(
        [
            0,
            20,
            40,
            60,
            80,
            100
        ]
    )
    ax.grid(
        linestyle="--",
        linewidth=0.5,
        alpha=0.25
    )
# ------------------------------------------------------------
# Information panel
# ------------------------------------------------------------
axes[
    7
].axis(
    "off"
)
legend_handles = [
    Line2D(
        [0],
        [0],
        color=CASE_COLORS[
            "VAMEI"
        ],
        linewidth=2.5,
        marker="o",
        markersize=5,
        label="VAMEI (2001)"
    ),
    Line2D(
        [0],
        [0],
        color=CASE_COLORS[
            "SENYAR"
        ],
        linewidth=2.5,
        marker="o",
        markersize=5,
        label="SENYAR (2025)"
    )
]
axes[
    7
].legend(
    handles=legend_handles,
    loc="upper left",
    title="Rare Event",
    fontsize=10,
    title_fontsize=11,
    frameon=True,
    framealpha=0.95
)
# Build the information panel from the actual observations rather than
# hard-coding genesis coordinates or life-cycle durations.
case_details = []
for sid, g in rare_tracks.groupby("SID"):
    g = g.sort_values("ISO_TIME")
    first = g.iloc[0]
    name = str(first["NAME"])
    latitude = float(first["LAT"])
    longitude = float(first["LON"])
    lat_hemisphere = "N" if latitude >= 0 else "S"
    lon_hemisphere = "E" if longitude >= 0 else "W"
    duration = float(g["Elapsed_Hours"].max())
    case_details.append(
        f"{name}: {abs(latitude):.1f}°{lat_hemisphere}, "
        f"{abs(longitude):.1f}°{lon_hemisphere} "
        f"| {duration:g} h"
    )
info_text = (
    "Genesis location | Duration:\n"
    + "\n".join(case_details)
    + "\n\nX-axis:\n"
    + "0% = genesis\n"
    + "100% = final recorded position\n\n"
    + "SST, TP, ω, ζ = ERA5\n"
    + "Wind, Pressure = Zoom Earth\n"
    + "Coriolis = calculated parameter"
)
axes[
    7
].text(
    0.02,
    0.54,
    info_text,
    transform=axes[
        7
    ].transAxes,
    va="top",
    fontsize=10.5,
    linespacing=1.4,
    bbox=dict(
        facecolor="white",
        edgecolor="gray",
        boxstyle="round,pad=0.5",
        alpha=0.95
    )
)
fig.suptitle(
    "Meteorological Evolution of Rare Equatorial Tropical Cyclones",
    fontsize=17,
    fontweight="bold",
    y=0.98
)
plt.tight_layout(
    rect=[
        0,
        0,
        1,
        0.95
    ]
)
plt.savefig(
    OUTPUT_G315,
    dpi=300,
    bbox_inches="tight",
    facecolor="white"
)
plt.show()
# ============================================================
# 23. FINAL OUTPUT
# ============================================================
print("\n" + "=" * 90)
print(
    "SECTION 3.1.3.4 ANALYSIS COMPLETED"
)
print("=" * 90)
print(
    "\nExcel:"
)
print(
    OUTPUT_EXCEL
)
print(
    "\nFigure 3.14:"
)
print(
    OUTPUT_G314
)
print(
    "\nFigure 3.15:"
)
print(
    OUTPUT_G315
)
print("=" * 90)
