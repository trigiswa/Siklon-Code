# ============================================================
# CHAPTER III - SECTION 3.1.1
# SPATIAL AND TEMPORAL DISTRIBUTION OF TROPICAL CYCLONES
#
# Input dataset: Siklon_Tropis_ERA5_Extracted.xlsx
# Study domain: 95°E–145°E, 15°S–25°N
# Event selection: Based on the genesis location
# Study period: 2001–2025
#
# Outputs:
# - Figure 3.1: Map of genesis locations by cyclone type
# - Figure 3.2: Annual number of tropical cyclone events
# - Figure 3.3: Monthly number of tropical cyclone events
# - Figure 3.4: Histogram of genesis latitude
# - Table 3.1: Number and percentage of events by zone
# - Table 3.2: Descriptive statistics of temporal distribution
# ============================================================

# 0. IMPORT LIBRARIES
import os
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import cartopy.crs as ccrs
import cartopy.feature as cfeature
from cartopy.mpl.gridliner import LONGITUDE_FORMATTER, LATITUDE_FORMATTER
from matplotlib.colors import ListedColormap, BoundaryNorm
from matplotlib.colorbar import ColorbarBase
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# 1. FILE SETTINGS
# Keep the original input filename and worksheet name to match your dataset.
FILE_PATH = "Siklon_Tropis_ERA5_Extracted.xlsx"
SHEET_NAME = "Sheet1"
OUTPUT_DIR = "output_3_1_1_english"
os.makedirs(OUTPUT_DIR, exist_ok=True)

# 2. STUDY DOMAIN
LON_MIN, LON_MAX = 95, 145
LAT_MIN, LAT_MAX = -15, 25
EQ_MIN, EQ_MAX = -5, 5  # Equatorial-zone boundaries
YEAR_START, YEAR_END = 2001, 2025

# 3. PLOT SETTINGS
plt.rcParams.update({
    "font.size": 11,
    "axes.labelsize": 12,
    "axes.titlesize": 14,
    "figure.dpi": 150,
    "savefig.dpi": 300,
    "font.family": "DejaVu Sans",
})

# 4. LOAD DATA
# Original field names are retained because they come from the input file.
df = pd.read_excel(FILE_PATH, sheet_name=SHEET_NAME)
print("=" * 65)
print("INITIAL DATA")
print("=" * 65)
print("Number of rows:", len(df))
print("Number of unique SIDs:", df["SID"].nunique())

# 5. DATA CLEANING
df["ISO_TIME"] = pd.to_datetime(df["ISO_TIME"], errors="coerce")
df["LAT"] = pd.to_numeric(df["LAT"], errors="coerce")
df["LON"] = pd.to_numeric(df["LON"], errors="coerce")
df["SEASON"] = pd.to_numeric(df["SEASON"], errors="coerce")

# Treat empty TYPE values and string representations of missing values as NA.
df["TYPE"] = df["TYPE"].astype("string").str.strip()
df["TYPE"] = df["TYPE"].replace({
    "": pd.NA,
    "nan": pd.NA,
    "None": pd.NA,
    "<NA>": pd.NA,
})

# Remove records missing key variables.
df = df.dropna(subset=["SID", "ISO_TIME", "LAT", "LON"]).copy()

# Retain valid geographic coordinates.
df = df[df["LAT"].between(-90, 90)].copy()
df = df[df["LON"].between(-180, 360)].copy()

# Remove duplicated SID–time records and sort by cyclone and timestamp.
df = df.drop_duplicates(subset=["SID", "ISO_TIME"], keep="first")
df = df.sort_values(["SID", "ISO_TIME"]).reset_index(drop=True)

print("\nAfter data cleaning:")
print("Number of rows:", len(df))
print("Number of unique SIDs:", df["SID"].nunique())

# 6. IDENTIFY THE FIRST RECORD FOR EACH SID
# After sorting by time, the first record per SID is taken as genesis,
# consistent with the original analysis.
genesis_all = df.drop_duplicates(subset="SID", keep="first").copy()
print("\nNumber of genesis events before spatial filtering:", len(genesis_all))

# 7. FILTER THE STUDY PERIOD (2001–2025)
genesis_all = genesis_all[
    genesis_all["SEASON"].between(YEAR_START, YEAR_END)
].copy()

# 8. FILTER EVENTS BY GENESIS LOCATION
genesis = genesis_all[
    genesis_all["LON"].between(LON_MIN, LON_MAX, inclusive="both")
    & genesis_all["LAT"].between(LAT_MIN, LAT_MAX, inclusive="both")
].copy()

# 9. CLASSIFY GEOGRAPHIC ZONES
def get_zone(lat):
    if EQ_MIN <= lat <= EQ_MAX:
        return "Equatorial"
    elif lat > EQ_MAX:
        return "Northern"
    else:
        return "Southern"


genesis["zone"] = genesis["LAT"].apply(get_zone)
genesis["month"] = genesis["ISO_TIME"].dt.month

# 10. QUALITY CONTROL
print("\n" + "=" * 65)
print("STUDY DATA AFTER GENESIS FILTERING")
print("=" * 65)
print("Number of events:", len(genesis))
print("\nNumber of events by zone:")
print(genesis["zone"].value_counts())
print("\nNumber of events by cyclone TYPE:")
print(genesis["TYPE"].value_counts(dropna=False))
print("=" * 65)

# Reference counts from the original dataset:
# Total = 392; Northern = 354; Equatorial = 2; Southern = 36.
if len(genesis) != 392:
    print(
        "\nWARNING:\n"
        "The total number of events is not 392.\n"
        "Please check the input file or preprocessing steps."
    )

# 11. TABLE 3.1 — NUMBER AND PERCENTAGE OF EVENTS BY ZONE
zone_order = ["Northern", "Equatorial", "Southern"]
zone_label = {
    "Northern": "North of the Equator",
    "Equatorial": "Equatorial Zone",
    "Southern": "South of the Equator",
}

table_31 = (
    genesis["zone"]
    .value_counts()
    .reindex(zone_order, fill_value=0)
    .rename_axis("Zone")
    .reset_index(name="Number of Events")
)
table_31["Percentage (%)"] = table_31["Number of Events"] / len(genesis) * 100
table_31["Zone"] = table_31["Zone"].map(zone_label)

total_31 = pd.DataFrame({
    "Zone": ["Total"],
    "Number of Events": [len(genesis)],
    "Percentage (%)": [100.0],
})
table_31 = pd.concat([table_31, total_31], ignore_index=True)

# 12. ANNUAL DISTRIBUTION
year_index = pd.Index(range(YEAR_START, YEAR_END + 1), name="Year")
yearly = (
    genesis.groupby(genesis["SEASON"].astype(int))
    .size()
    .reindex(year_index, fill_value=0)
    .rename("Number of Events")
    .reset_index()
)
yearly_mean = yearly["Number of Events"].mean()
max_year_count = yearly["Number of Events"].max()
max_years = yearly.loc[
    yearly["Number of Events"] == max_year_count, "Year"
].tolist()
min_year_count = yearly["Number of Events"].min()
min_years = yearly.loc[
    yearly["Number of Events"] == min_year_count, "Year"
].tolist()

# 13. MONTHLY DISTRIBUTION
month_index = pd.Index(range(1, 13), name="Month")
monthly = (
    genesis.groupby("month")
    .size()
    .reindex(month_index, fill_value=0)
    .rename("Number of Events")
    .reset_index()
)
month_names = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]
month_abbreviations = [
    "Jan", "Feb", "Mar", "Apr", "May", "Jun",
    "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
]
monthly["Month Name"] = month_names
max_month_count = monthly["Number of Events"].max()
max_month_row = monthly.loc[monthly["Number of Events"].idxmax()]
max_month = max_month_row["Month Name"]
min_month_count = monthly["Number of Events"].min()
min_month_row = monthly.loc[monthly["Number of Events"].idxmin()]
min_month = min_month_row["Month Name"]

# 14. SEASONAL DISTRIBUTION
def get_season_group(month):
    if month in [12, 1, 2]:
        return "DJF"
    elif month in [3, 4, 5]:
        return "MAM"
    elif month in [6, 7, 8]:
        return "JJA"
    else:
        return "SON"


genesis["Seasonal Period"] = genesis["month"].apply(get_season_group)
season_order = ["DJF", "MAM", "JJA", "SON"]
season_months = {
    "DJF": "December–January–February",
    "MAM": "March–April–May",
    "JJA": "June–July–August",
    "SON": "September–October–November",
}
seasonal = (
    genesis["Seasonal Period"]
    .value_counts()
    .reindex(season_order, fill_value=0)
    .rename_axis("Period")
    .reset_index(name="Number of Events")
)
seasonal["Months"] = seasonal["Period"].map(season_months)
seasonal["Percentage (%)"] = seasonal["Number of Events"] / len(genesis) * 100
seasonal = seasonal[["Period", "Months", "Number of Events", "Percentage (%)"]]

dominant_season_row = seasonal.loc[seasonal["Number of Events"].idxmax()]
dominant_season = dominant_season_row["Period"]
dominant_season_count = dominant_season_row["Number of Events"]
dominant_season_pct = dominant_season_row["Percentage (%)"]

# 15. TABLE 3.2 — DESCRIPTIVE STATISTICS OF TEMPORAL DISTRIBUTION
table_32 = pd.DataFrame({
    "Parameter": [
        "Total number of events",
        "Mean number of events per year",
        "Year(s) with the highest number of events",
        "Highest annual event count",
        "Year(s) with the lowest number of events",
        "Lowest annual event count",
        "Month with the highest number of events",
        "Highest monthly event count",
        "Month with the lowest number of events",
        "Lowest monthly event count",
        "Dominant seasonal period",
        "Number of events in the dominant seasonal period",
        "Percentage in the dominant seasonal period (%)",
    ],
    "Value": [
        len(genesis),
        round(yearly_mean, 2),
        ", ".join(map(str, max_years)),
        int(max_year_count),
        ", ".join(map(str, min_years)),
        int(min_year_count),
        max_month,
        int(max_month_count),
        min_month,
        int(min_month_count),
        dominant_season,
        int(dominant_season_count),
        round(dominant_season_pct, 2),
    ],
})

# 16. CYCLONE TYPE DISTRIBUTION
type_order = ["S", "M", "L", "D", "C"]
type_distribution = (
    genesis["TYPE"]
    .value_counts(dropna=False)
    .rename_axis("TYPE")
    .reset_index(name="Number of Genesis Events")
)
type_distribution["TYPE"] = type_distribution["TYPE"].fillna("Missing/NA")
type_distribution["Percentage (%)"] = (
    type_distribution["Number of Genesis Events"] / len(genesis) * 100
)

# 17. SAVE TABLES AND SUPPORTING DATA TO EXCEL
OUTPUT_EXCEL = os.path.join(OUTPUT_DIR, "Chapter_III_Section_3_1_1_Tables.xlsx")
with pd.ExcelWriter(OUTPUT_EXCEL, engine="openpyxl") as writer:
    table_31.round(1).to_excel(writer, sheet_name="Table 3.1", index=False)
    table_32.to_excel(writer, sheet_name="Table 3.2", index=False)
    seasonal.round(2).to_excel(writer, sheet_name="Seasonal Data", index=False)
    yearly.to_excel(writer, sheet_name="Annual Data", index=False)
    monthly.to_excel(writer, sheet_name="Monthly Data", index=False)
    type_distribution.round(2).to_excel(writer, sheet_name="TYPE Distribution", index=False)
    genesis.to_excel(writer, sheet_name="Genesis Data", index=False)

# 18. FORMAT THE EXCEL WORKBOOK
wb = load_workbook(OUTPUT_EXCEL)
header_fill = PatternFill("solid", fgColor="D9EAF7")
header_font = Font(bold=True)
thin_border = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)

for ws in wb.worksheets:
    ws.freeze_panes = "A2"

    # Format the header row.
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Add thin borders and vertical alignment to all cells.
    for row in ws.iter_rows():
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    # Adjust column widths according to the longest cell content.
    for column_cells in ws.columns:
        column_letter = column_cells[0].column_letter
        max_length = max(
            (len(str(cell.value)) for cell in column_cells if cell.value is not None),
            default=0,
        )
        ws.column_dimensions[column_letter].width = min(max_length + 2, 42)

wb.save(OUTPUT_EXCEL)
print("\nExcel file saved successfully:")
print(OUTPUT_EXCEL)

# ============================================================
# FIGURE 3.1 — SPATIAL DISTRIBUTION OF GENESIS BY CYCLONE TYPE
# ============================================================
OUTPUT_G31 = os.path.join(OUTPUT_DIR, "Figure_3_1_Genesis_Locations_by_Type.png")

# 19. CYCLONE TYPE COLORS
type_colors = {
    "S": "#e74c3c",
    "M": "#f39c12",
    "L": "#f7dc6f",
    "D": "#5dade2",
    "C": "#1a5276",
}

# 20. CREATE MAP FIGURE
fig = plt.figure(figsize=(12, 9))
ax = plt.axes(projection=ccrs.PlateCarree())
ax.set_extent([LON_MIN, LON_MAX, LAT_MIN, LAT_MAX], crs=ccrs.PlateCarree())

# 21. BASEMAP
ax.add_feature(
    cfeature.OCEAN.with_scale("10m"), facecolor="#d5dbdb", zorder=0
)
ax.add_feature(
    cfeature.LAND.with_scale("10m"), facecolor="#f8f9f9", zorder=0
)
ax.add_feature(
    cfeature.COASTLINE.with_scale("10m"),
    linewidth=0.7, edgecolor="#2c3e50", zorder=2,
)
ax.add_feature(
    cfeature.BORDERS.with_scale("10m"),
    linestyle=":", linewidth=0.5, edgecolor="#7f8c8d", zorder=2,
)

# 22. GRIDLINES AND COORDINATE LABELS
gl = ax.gridlines(
    draw_labels=True, linewidth=0.35, color="gray",
    alpha=0.45, linestyle="--",
)
gl.top_labels = False
gl.right_labels = False
gl.xformatter = LONGITUDE_FORMATTER
gl.yformatter = LATITUDE_FORMATTER

# 23. PLOT GENESIS LOCATIONS BY CYCLONE TYPE
for cyclone_type in type_order:
    subset = genesis[genesis["TYPE"] == cyclone_type]
    if len(subset) > 0:
        ax.scatter(
            subset["LON"], subset["LAT"],
            color=type_colors[cyclone_type], s=40, alpha=0.90,
            edgecolors="#222222", linewidths=0.35,
            transform=ccrs.PlateCarree(), zorder=5,
        )

# 24. PLOT RECORDS WITH MISSING TYPE
# These events are displayed as grey crosses without legend entries.
missing_type = genesis[genesis["TYPE"].isna()]
if len(missing_type) > 0:
    ax.scatter(
        missing_type["LON"], missing_type["LAT"],
        color="#808080", s=46, marker="x", linewidths=1.3,
        transform=ccrs.PlateCarree(), zorder=6,
    )

# 25. EQUATOR LINE
ax.plot(
    [LON_MIN, LON_MAX], [0, 0],
    color="#2c3e50", linestyle="--", linewidth=1.0,
    alpha=0.8, transform=ccrs.PlateCarree(), zorder=3,
)

# 26. FIGURE TITLE
ax.set_title(
    "Spatial Distribution of Tropical Cyclone Genesis in Southeast Asia",
    fontsize=14, pad=12,
)

# 27. COLORBAR FOR CYCLONE TYPE (S AT THE TOP, C AT THE BOTTOM)
cax = fig.add_axes([0.92, 0.22, 0.025, 0.55])
cmap = ListedColormap([type_colors[t] for t in reversed(type_order)])
bounds = np.arange(len(type_order) + 1)
norm = BoundaryNorm(bounds, cmap.N)
cb = ColorbarBase(cax, cmap=cmap, norm=norm, orientation="vertical")
cb.set_ticks(np.arange(len(type_order)) + 0.5)
cb.set_ticklabels(list(reversed(type_order)))
cb.set_label("Cyclone Type", fontsize=11, labelpad=8)
cb.ax.tick_params(labelsize=10)

# 28. SAVE FIGURE 3.1
plt.tight_layout(rect=[0, 0, 0.90, 1])
plt.savefig(OUTPUT_G31, dpi=300, bbox_inches="tight", facecolor="white")
plt.show()
print("\nFigure 3.1 saved:")
print(OUTPUT_G31)

# ============================================================
# FIGURE 3.2 — ANNUAL NUMBER OF TROPICAL CYCLONE EVENTS
# ============================================================
OUTPUT_G32 = os.path.join(OUTPUT_DIR, "Figure_3_2_Annual_Event_Count.png")
fig, ax = plt.subplots(figsize=(12, 5))
bars = ax.bar(
    yearly["Year"], yearly["Number of Events"],
    color="#1f77b4", edgecolor="black", width=0.7,
)
ax.set_xlabel("Year")
ax.set_ylabel("Number of Events")
ax.set_title("Annual Number of Tropical Cyclone Events (2001–2025)")
ax.set_xticks(yearly["Year"])
ax.tick_params(axis="x", rotation=45)

# Display the count above each bar.
for bar in bars:
    height = bar.get_height()
    ax.annotate(
        f"{int(height)}",
        xy=(bar.get_x() + bar.get_width() / 2, height),
        xytext=(0, 3), textcoords="offset points",
        ha="center", va="bottom", fontsize=8,
    )

# Annual mean reference line.
ax.axhline(
    yearly_mean, color="red", linestyle="--", linewidth=1.2,
    label=f"Annual Mean ({yearly_mean:.2f})",
)
ax.legend()
plt.tight_layout()
plt.savefig(OUTPUT_G32, dpi=300, bbox_inches="tight")
plt.show()

# ============================================================
# FIGURE 3.3 — MONTHLY NUMBER OF TROPICAL CYCLONE EVENTS
# ============================================================
OUTPUT_G33 = os.path.join(OUTPUT_DIR, "Figure_3_3_Monthly_Event_Count.png")
fig, ax = plt.subplots(figsize=(10, 5))
bars = ax.bar(
    monthly["Month"], monthly["Number of Events"],
    color="#2ca02c", edgecolor="black", width=0.7,
)
ax.set_xlabel("Month")
ax.set_ylabel("Number of Events")
ax.set_title("Monthly Number of Tropical Cyclone Events (2001–2025)")
ax.set_xticks(range(1, 13))
ax.set_xticklabels(month_abbreviations)

# Display the count above each bar.
for bar in bars:
    height = bar.get_height()
    ax.annotate(
        f"{int(height)}",
        xy=(bar.get_x() + bar.get_width() / 2, height),
        xytext=(0, 3), textcoords="offset points",
        ha="center", va="bottom", fontsize=9,
    )

plt.tight_layout()
plt.savefig(OUTPUT_G33, dpi=300, bbox_inches="tight")
plt.show()

# ============================================================
# FIGURE 3.4 — HISTOGRAM OF TROPICAL CYCLONE GENESIS LATITUDE
# ============================================================
OUTPUT_G34 = os.path.join(OUTPUT_DIR, "Figure_3_4_Genesis_Latitude_Histogram.png")
fig, ax = plt.subplots(figsize=(9, 5))
ax.hist(
    genesis["LAT"], bins=25,
    color="#1f77b4", edgecolor="black", alpha=0.70,
)

# Equatorial-zone limits and the equator.
ax.axvline(5, color="red", linestyle="--", linewidth=1.5, label="+5° Boundary")
ax.axvline(-5, color="red", linestyle="--", linewidth=1.5, label="−5° Boundary")
ax.axvline(0, color="black", linestyle="-", linewidth=1.0, alpha=0.7, label="Equator")
ax.set_xlabel("Genesis Latitude (°)")
ax.set_ylabel("Frequency")
ax.set_title("Distribution of Tropical Cyclone Genesis Latitude (2001–2025)")
ax.legend()
plt.tight_layout()
plt.savefig(OUTPUT_G34, dpi=300, bbox_inches="tight")
plt.show()

# 29. PRINT FINAL SUMMARY
print("\n" + "=" * 70)
print("FINAL SUMMARY — SECTION 3.1.1")
print("=" * 70)
print(f"Total number of events: {len(genesis)}")
print(f"Mean events per year: {yearly_mean:.2f}")
print(f"Year(s) with highest count: {max_years} ({int(max_year_count)} events)")
print(f"Year(s) with lowest count: {min_years} ({int(min_year_count)} events)")
print(f"Month with highest count: {max_month} ({int(max_month_count)} events)")
print(f"Month with lowest count: {min_month} ({int(min_month_count)} events)")
print(
    f"Dominant seasonal period: {dominant_season} "
    f"({int(dominant_season_count)} events; {dominant_season_pct:.2f}%)"
)
print("\nTABLE 3.1")
print(table_31.round(1).to_string(index=False))
print("\nSEASONAL DISTRIBUTION")
print(seasonal.round(2).to_string(index=False))
print("\nCYCLONE TYPE DISTRIBUTION")
print(type_distribution.round(2).to_string(index=False))
print("\nAll files have been saved to:")
print(OUTPUT_DIR)
print("=" * 70)
