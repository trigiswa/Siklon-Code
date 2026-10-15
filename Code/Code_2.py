# ============================================================

# CHAPTER III - 3.1.2.2
# DETERMINATION OF THE OPTIMAL NUMBER OF CLUSTERS
#
# Methods:
# 1. Select SIDs based on genesis within the study domain
# 2. Resample trajectories
# 3. Dynamic Time Warping (DTW)
# 4. K-Medoids
# 5. Evaluate Silhouette Scores for k = 2 to 8
#
# Output:
# - Figure 3.6: Silhouette Score
# - Table 3.4: Clustering evaluation
# - Optimal cluster assignment for each SID
# - List of medoids for each cluster
# - DTW distance matrix (.npy) for subsequent analyses

# ============================================================

# ============================================================

# 0. IMPORT LIBRARY

# ============================================================

import os
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from numba import njit, prange
from sklearn.metrics import silhouette_score
from openpyxl import load_workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)

# ============================================================

# 1. FILE SETTINGS

# ============================================================

FILE_PATH = "Siklon_Tropis_ERA5_Extracted.xlsx"
SHEET_NAME = "Sheet1"
OUTPUT_DIR = "output_3_1_2"
os.makedirs(
    OUTPUT_DIR,
    exist_ok=True
)
OUTPUT_EXCEL = os.path.join(
    OUTPUT_DIR,
    "Clustering_Results_3_1_2_2.xlsx"
)
OUTPUT_G36 = os.path.join(
    OUTPUT_DIR,
    "Figure_3.6_Silhouette_Score.png"
)
OUTPUT_DTW = os.path.join(
    OUTPUT_DIR,
    "DTW_Distance_Matrix.npy"
)
OUTPUT_SID = os.path.join(
    OUTPUT_DIR,
    "DTW_SID_Order.npy"
)

# ============================================================

# 2. STUDY DOMAIN

# ============================================================

LON_MIN = 95
LON_MAX = 145
LAT_MIN = -15
LAT_MAX = 25
EQ_MIN = -5
EQ_MAX = 5
YEAR_START = 2001
YEAR_END = 2025

# ============================================================

# 3. CLUSTERING SETTINGS

# ============================================================

K_MIN = 2
K_MAX = 8
RANDOM_STATE = 42

# Multiple initializations are used so that the K-Medoids solution
# does not depend excessively on a single initial set of medoids.
N_INIT = 30
MAX_ITER = 100

# ============================================================

# 4. VISUAL SETTINGS

# ============================================================

plt.rcParams.update({
    "font.size": 11,
    "axes.labelsize": 12,
    "axes.titlesize": 14,
    "figure.dpi": 150,
    "savefig.dpi": 300,
    "font.family": "DejaVu Sans"
})

# ============================================================

# 5. LOAD DATA

# ============================================================

df = pd.read_excel(
    FILE_PATH,
    sheet_name=SHEET_NAME
)

# ============================================================

# 6. CLEANING

# ============================================================

df["ISO_TIME"] = pd.to_datetime(
    df["ISO_TIME"],
    errors="coerce"
)
df["LAT"] = pd.to_numeric(
    df["LAT"],
    errors="coerce"
)
df["LON"] = pd.to_numeric(
    df["LON"],
    errors="coerce"
)
df["SEASON"] = pd.to_numeric(
    df["SEASON"],
    errors="coerce"
)
df = df.dropna(
    subset=[
        "SID",
        "ISO_TIME",
        "LAT",
        "LON"
    ]
).copy()
df = df[
    df["LAT"].between(-90, 90)
].copy()
df = df[
    df["LON"].between(-180, 360)
].copy()
df = df.drop_duplicates(
    subset=[
        "SID",
        "ISO_TIME"
    ],
    keep="first"
)
df = df.sort_values(
    [
        "SID",
        "ISO_TIME"
    ]
).reset_index(
    drop=True
)

# ============================================================

# 7. IDENTIFY GENESIS

# ============================================================

genesis_all = (
    df.drop_duplicates(
        subset="SID",
        keep="first"
    )
    .copy()
)
genesis_all = genesis_all[
    genesis_all["SEASON"].between(
        YEAR_START,
        YEAR_END
    )
].copy()

# ============================================================

# 8. FILTER BY STUDY DOMAIN

# ============================================================

genesis = genesis_all[
    genesis_all["LON"].between(
        LON_MIN,
        LON_MAX
    )
    &
    genesis_all["LAT"].between(
        LAT_MIN,
        LAT_MAX
    )
].copy()

# ============================================================

# 9. GENESIS ZONES

# ============================================================

def get_zone(lat):
    if EQ_MIN <= lat <= EQ_MAX:
        return "Equatorial"
    elif lat > EQ_MAX:
        return "Northern Hemisphere"
    else:
        return "Southern Hemisphere"
genesis["Genesis_Zone"] = (
    genesis["LAT"]
    .apply(get_zone)
)

# ============================================================

# 10. EXTRACT COMPLETE TRACKS FOR SELECTED SIDS

# ============================================================

selected_sids = (
    genesis["SID"]
    .unique()
)
track = df[
    df["SID"].isin(
        selected_sids
    )
].copy()
print("=" * 70)
print("DATA CLUSTERING")
print("=" * 70)
print(
    "Number of storm IDs (SID):",
    len(selected_sids)
)
print(
    "Number of track points:",
    len(track)
)
print("=" * 70)

# Quality control
if len(selected_sids) != 392:
    print(
        "WARNING: "
        "The number of SIDs is not 392."
    )

# ============================================================

# 11. NUMBER OF POINTS IN EACH TRAJECTORY

# ============================================================

points_per_track = (
    track
    .groupby(
        "SID"
    )
    .size()
)
print(
    "\nTrack-point count statistics:"
)
print(
    points_per_track.describe()
)

# ============================================================

# 12. NUMBER OF RESAMPLING POINTS

# ============================================================

#
# Use the median number of original points to ensure that:
# - short trajectories are not excessively upsampled
# - long trajectories do not lose excessive information
#
# For the current dataset, the median is 44.

# ============================================================

N_RESAMPLE = int(
    round(
        points_per_track.median()
    )
)
print(
    "\nNumber of resampled points:",
    N_RESAMPLE
)

# ============================================================

# 13. TRAJECTORY RESAMPLING FUNCTION

# ============================================================

#
# Time is normalized for each track:
#
# genesis = 0
# termination = 1
#
# Latitude and longitude are then linearly interpolated.

# ============================================================

def resample_trajectory(
    group,
    n_points
):
    group = group.sort_values(
        "ISO_TIME"
    ).copy()
    lat = (
        group["LAT"]
        .to_numpy(
            dtype=float
        )
    )
    lon = (
        group["LON"]
        .to_numpy(
            dtype=float
        )
    )

    # Relative time in seconds
    relative_time = (
        group["ISO_TIME"]
        -
        group[
            "ISO_TIME"
        ].iloc[0]
    ).dt.total_seconds().to_numpy(
        dtype=float
    )

    # For tracks with only one position
    # or zero duration
    if (
        len(group) == 1
        or
        relative_time[-1] == 0
    ):
        lat_new = np.full(
            n_points,
            lat[0]
        )
        lon_new = np.full(
            n_points,
            lon[0]
        )
        return np.column_stack(
            [
                lat_new,
                lon_new
            ]
        )

    # Normalize time to [0, 1]
    normalized_time = (
        relative_time
        /
        relative_time[-1]
    )
    target_time = np.linspace(
        0,
        1,
        n_points
    )

    # Interpolate latitude
    lat_new = np.interp(
        target_time,
        normalized_time,
        lat
    )

    # Interpolate longitude
    lon_new = np.interp(
        target_time,
        normalized_time,
        lon
    )
    return np.column_stack(
        [
            lat_new,
            lon_new
        ]
    )

# ============================================================

# 14. RESAMPLE ALL TRAJECTORIES

# ============================================================

sid_order = []
resampled_tracks = []
for sid, group in track.groupby(
    "SID"
):
    sid_order.append(
        sid
    )
    resampled_tracks.append(
        resample_trajectory(
            group,
            N_RESAMPLE
        )
    )
resampled_tracks = np.stack(
    resampled_tracks
)
sid_order = np.array(
    sid_order
)
print(
    "\nTrajectory array shape:"
)
print(
    resampled_tracks.shape
)

# ============================================================

# 15. HAVERSINE DISTANCE

# ============================================================

#
# Used as the local distance measure within DTW.
# Distance is measured in kilometers.

# ============================================================

EARTH_RADIUS_KM = 6371.0088
@njit
def haversine_numba(
    lat1,
    lon1,
    lat2,
    lon2
):
    p1 = (
        lat1
        *
        np.pi
        /
        180.0
    )
    p2 = (
        lat2
        *
        np.pi
        /
        180.0
    )
    dlat = (
        (
            lat2
            -
            lat1
        )
        *
        np.pi
        /
        180.0
    )
    dlon = (
        (
            lon2
            -
            lon1
        )
        *
        np.pi
        /
        180.0
    )
    a = (
        np.sin(
            dlat / 2.0
        ) ** 2
        +
        np.cos(
            p1
        )
        *
        np.cos(
            p2
        )
        *
        np.sin(
            dlon / 2.0
        ) ** 2
    )
    c = (
        2.0
        *
        np.arctan2(
            np.sqrt(a),
            np.sqrt(
                1.0 - a
            )
        )
    )
    return (
        EARTH_RADIUS_KM
        *
        c
    )

# ============================================================

# 16. DYNAMIC TIME WARPING

# ============================================================

@njit
def dtw_distance(
    track_a,
    track_b
):
    n = track_a.shape[0]
    m = track_b.shape[0]

    # Only two rows are stored
    # to reduce memory usage.
    previous = np.full(
        m + 1,
        1e30
    )
    current = np.full(
        m + 1,
        1e30
    )
    previous[0] = 0.0
    for i in range(
        1,
        n + 1
    ):
        current[0] = 1e30
        for j in range(
            1,
            m + 1
        ):
            cost = haversine_numba(
                track_a[
                    i - 1,
                    0
                ],
                track_a[
                    i - 1,
                    1
                ],
                track_b[
                    j - 1,
                    0
                ],
                track_b[
                    j - 1,
                    1
                ]
            )
            minimum_previous = previous[
                j
            ]
            if (
                current[
                    j - 1
                ]
                <
                minimum_previous
            ):
                minimum_previous = current[
                    j - 1
                ]
            if (
                previous[
                    j - 1
                ]
                <
                minimum_previous
            ):
                minimum_previous = previous[
                    j - 1
                ]
            current[
                j
            ] = (
                cost
                +
                minimum_previous
            )
        temporary = previous
        previous = current
        current = temporary
    return previous[m]

# ============================================================

# 17. DTW DISTANCE MATRIX

# ============================================================

@njit(
    parallel=True
)
def pairwise_dtw_matrix(
    trajectories
):
    n_track = (
        trajectories.shape[0]
    )
    distance_matrix = np.zeros(
        (
            n_track,
            n_track
        ),
        dtype=np.float64
    )
    for i in prange(
        n_track
    ):
        for j in range(
            i + 1,
            n_track
        ):
            distance = dtw_distance(
                trajectories[i],
                trajectories[j]
            )
            distance_matrix[
                i,
                j
            ] = distance
            distance_matrix[
                j,
                i
            ] = distance
    return distance_matrix
print(
    "\nComputing the DTW distance matrix..."
)

# The first call also triggers
# compilation of the Numba function.
distance_matrix = pairwise_dtw_matrix(
    resampled_tracks
)
print(
    "DTW distance matrix calculation completed."
)
print(
    "Matrix shape:",
    distance_matrix.shape
)

# Save the results to avoid recalculation
np.save(
    OUTPUT_DTW,
    distance_matrix
)
np.save(
    OUTPUT_SID,
    sid_order
)

# ============================================================

# 18. K-MEDOIDS++ INITIALIZATION

# ============================================================

def initialize_kmedoids_pp(
    distance_matrix,
    k,
    rng
):
    n = (
        distance_matrix.shape[0]
    )

    # Select the first medoid randomly
    medoids = [
        int(
            rng.integers(
                n
            )
        )
    ]
    while len(
        medoids
    ) < k:
        nearest_distance = np.min(
            distance_matrix[
                :,
                medoids
            ],
            axis=1
        )

        # Do not select the same medoid again
        nearest_distance[
            medoids
        ] = 0.0
        probabilities = (
            nearest_distance
            **
            2
        )
        total = (
            probabilities.sum()
        )
        if total <= 0:
            candidates = np.setdiff1d(
                np.arange(
                    n
                ),
                np.array(
                    medoids
                )
            )
            new_medoid = int(
                rng.choice(
                    candidates
                )
            )
        else:
            probabilities = (
                probabilities
                /
                total
            )
            new_medoid = int(
                rng.choice(
                    n,
                    p=probabilities
                )
            )
        if (
            new_medoid
            not in medoids
        ):
            medoids.append(
                new_medoid
            )
    return np.array(
        medoids,
        dtype=int
    )

# ============================================================

# 19. K-MEDOIDS

# ============================================================

def kmedoids_precomputed(
    distance_matrix,
    k,
    random_state=42,
    n_init=30,
    max_iter=100
):
    n = (
        distance_matrix.shape[0]
    )
    base_rng = np.random.default_rng(
        random_state
    )
    seeds = base_rng.integers(
        0,
        2**32 - 1,
        size=n_init,
        dtype=np.uint32
    )
    best_result = None
    for seed in seeds:
        rng = np.random.default_rng(
            int(seed)
        )
        medoids = initialize_kmedoids_pp(
            distance_matrix,
            k,
            rng
        )
        labels = None
        for iteration in range(
            max_iter
        ):

            # ----------------------------------------------
            # Assignment
            # ----------------------------------------------
            labels = np.argmin(
                distance_matrix[
                    :,
                    medoids
                ],
                axis=1
            )
            new_medoids = (
                medoids.copy()
            )

            # ----------------------------------------------
            # Update the medoid of each cluster
            # ----------------------------------------------
            for cluster_id in range(
                k
            ):
                members = np.where(
                    labels
                    ==
                    cluster_id
                )[0]

                # Empty cluster
                if len(
                    members
                ) == 0:
                    nearest = np.min(
                        distance_matrix[
                            :,
                            medoids
                        ],
                        axis=1
                    )
                    candidates = np.argsort(
                        nearest
                    )[::-1]
                    for candidate in candidates:
                        if (
                            candidate
                            not in
                            new_medoids
                        ):
                            new_medoids[
                                cluster_id
                            ] = candidate
                            break
                else:
                    sub_matrix = (
                        distance_matrix[
                            np.ix_(
                                members,
                                members
                            )
                        ]
                    )
                    total_distance = (
                        sub_matrix.sum(
                            axis=1
                        )
                    )
                    best_member = (
                        members[
                            np.argmin(
                                total_distance
                            )
                        ]
                    )
                    new_medoids[
                        cluster_id
                    ] = best_member

            # ----------------------------------------------
            # Convergence
            # ----------------------------------------------
            if np.array_equal(
                np.sort(
                    new_medoids
                ),
                np.sort(
                    medoids
                )
            ):
                medoids = new_medoids
                labels = np.argmin(
                    distance_matrix[
                        :,
                        medoids
                    ],
                    axis=1
                )
                break
            medoids = new_medoids

        # ----------------------------------------------
        # Objective:
        # sum of distances from each observation
        # to its cluster medoid
        # ----------------------------------------------
        inertia = np.sum(
            distance_matrix[
                np.arange(
                    n
                ),
                medoids[
                    labels
                ]
            ]
        )
        if (
            best_result is None
            or
            inertia
            <
            best_result[
                "inertia"
            ]
        ):
            best_result = {
                "medoids":
                    medoids.copy(),
                "labels":
                    labels.copy(),
                "inertia":
                    inertia,
                "n_iter":
                    iteration + 1
            }
    return best_result

# ============================================================

# 20. EVALUATE k FROM 2 TO 8

# ============================================================

evaluation_records = []
cluster_results = {}
print(
    "\nComputing K-Medoids clustering and Silhouette Scores..."
)
for k in range(
    K_MIN,
    K_MAX + 1
):
    result = kmedoids_precomputed(
        distance_matrix,
        k,
        random_state=RANDOM_STATE,
        n_init=N_INIT,
        max_iter=MAX_ITER
    )
    silhouette = silhouette_score(
        distance_matrix,
        result[
            "labels"
        ],
        metric="precomputed"
    )
    cluster_results[
        k
    ] = result
    evaluation_records.append({
        "Number of Clusters (k)":
            k,
        "Silhouette Score":
            silhouette,
        "Total Within-Cluster Distance":
            result[
                "inertia"
            ]
    })
    print(
        f"k = {k} | "
        f"Silhouette = "
        f"{silhouette:.4f}"
    )
evaluation_df = pd.DataFrame(
    evaluation_records
)

# ============================================================

# 21. SELECT THE OPTIMAL k

# ============================================================

optimal_row = evaluation_df.loc[
    evaluation_df[
        "Silhouette Score"
    ].idxmax()
]
OPTIMAL_K = int(
    optimal_row[
        "Number of Clusters (k)"
    ]
)
OPTIMAL_SILHOUETTE = float(
    optimal_row[
        "Silhouette Score"
    ]
)
print(
    "\nOptimal number of clusters:",
    OPTIMAL_K
)
print(
    "Highest Silhouette Score:",
    round(
        OPTIMAL_SILHOUETTE,
        4
    )
)

# ============================================================

# 22. RETRIEVE OPTIMAL CLUSTERING RESULT

# ============================================================

optimal_result = (
    cluster_results[
        OPTIMAL_K
    ]
)
raw_labels = (
    optimal_result[
        "labels"
    ]
)
raw_medoids = (
    optimal_result[
        "medoids"
    ]
)

# ============================================================

# 23. METADATA FOR EACH SID

# ============================================================

metadata = (
    genesis[
        [
            "SID",
            "NAME",
            "SEASON",
            "LAT",
            "LON",
            "Genesis_Zone"
        ]
    ]
    .copy()
)
metadata = metadata.rename(
    columns={
        "LAT":
            "Genesis_Lat",
        "LON":
            "Genesis_Lon"
    }
)

# Align metadata order exactly with the DTW matrix
metadata = (
    metadata
    .set_index(
        "SID"
    )
    .loc[
        sid_order
    ]
    .reset_index()
)
metadata[
    "Raw_Cluster"
] = (
    raw_labels
)

# ============================================================

# 24. RENUMBER CLUSTERS FOR PRESENTATION

# ============================================================

#
# Cluster labels assigned by the algorithm are arbitrary.
#
# For a consistent presentation,
# clusters are sorted by median genesis latitude
# from north to south.

# ============================================================

cluster_median_lat = (
    metadata
    .groupby(
        "Raw_Cluster"
    )[
        "Genesis_Lat"
    ]
    .median()
    .sort_values(
        ascending=False
    )
)
cluster_mapping = {
    raw_cluster:
        new_cluster
    for new_cluster,
        raw_cluster
    in enumerate(
        cluster_median_lat.index,
        start=1
    )
}
metadata[
    "Cluster"
] = (
    metadata[
        "Raw_Cluster"
    ]
    .map(
        cluster_mapping
    )
)

# ============================================================

# 25. MEDOID DATA

# ============================================================

medoid_records = []
for raw_cluster_id in range(
    OPTIMAL_K
):
    medoid_index = (
        raw_medoids[
            raw_cluster_id
        ]
    )
    medoid_sid = (
        sid_order[
            medoid_index
        ]
    )
    row = metadata[
        metadata[
            "SID"
        ]
        ==
        medoid_sid
    ].iloc[0]
    medoid_records.append({
        "Cluster":
            cluster_mapping[
                raw_cluster_id
            ],
        "Medoid SID":
            row[
                "SID"
            ],
        "Cyclone Name":
            row[
                "NAME"
            ],
        "Genesis Latitude":
            row[
                "Genesis_Lat"
            ],
        "Genesis Longitude":
            row[
                "Genesis_Lon"
            ],
        "Genesis Zone":
            row[
                "Genesis_Zone"
            ]
    })
medoid_df = (
    pd.DataFrame(
        medoid_records
    )
    .sort_values(
        "Cluster"
    )
    .reset_index(
        drop=True
    )
)

# ============================================================

# 26. CLUSTER MEMBERSHIP COUNTS

# ============================================================

cluster_size = (
    metadata[
        "Cluster"
    ]
    .value_counts()
    .sort_index()
    .rename_axis(
        "Cluster"
    )
    .reset_index(
        name="Number of Members"
    )
)
cluster_size[
    "Percentage (%)"
] = (
    cluster_size[
        "Number of Members"
    ]
    /
    len(
        metadata
    )
    *
    100
)

# ============================================================

# 27. TABLE 3.4

# ============================================================

table_34 = (
    evaluation_df[
        [
            "Number of Clusters (k)",
            "Silhouette Score"
        ]
    ]
    .copy()
)

# ============================================================

# 28. METHOD SUMMARY

# ============================================================

method_summary = pd.DataFrame({
    "Parameter": [
        "Number of trajectories",
        "Median number of original track points",
        "Number of resampled points",
        "Range of evaluated k values",
        "Number of K-Medoids initializations",
        "Random state",
        "Optimal number of clusters",
        "Optimal Silhouette Score"
    ],
    "Value": [
        len(
            sid_order
        ),
        float(
            points_per_track.median()
        ),
        N_RESAMPLE,
        f"{K_MIN}–{K_MAX}",
        N_INIT,
        RANDOM_STATE,
        OPTIMAL_K,
        round(
            OPTIMAL_SILHOUETTE,
            4
        )
    ]
})

# ============================================================

# 29. SAVE TO EXCEL

# ============================================================

with pd.ExcelWriter(
    OUTPUT_EXCEL,
    engine="openpyxl"
) as writer:

    # Table included in the thesis
    table_34.round(
        4
    ).to_excel(
        writer,
        sheet_name="Table 3.4",
        index=False
    )

    # Method summary
    method_summary.to_excel(
        writer,
        sheet_name="Method Summary",
        index=False
    )

    # Cluster assignments for analysis in Section 3.1.2.3
    metadata[
        [
            "SID",
            "NAME",
            "SEASON",
            "Genesis_Lat",
            "Genesis_Lon",
            "Genesis_Zone",
            "Cluster"
        ]
    ].to_excel(
        writer,
        sheet_name="Cluster Assignments",
        index=False
    )

    # Cluster membership counts
    cluster_size.round(
        2
    ).to_excel(
        writer,
        sheet_name="Cluster Sizes",
        index=False
    )

    # Medoids
    medoid_df.to_excel(
        writer,
        sheet_name="Cluster Medoids",
        index=False
    )

# ============================================================

# 30. FORMAT EXCEL WORKBOOK

# ============================================================

wb = load_workbook(
    OUTPUT_EXCEL
)
header_fill = PatternFill(
    "solid",
    fgColor="D9EAF7"
)
header_font = Font(
    bold=True
)
thin_border = Border(
    left=Side(
        style="thin",
        color="BFBFBF"
    ),
    right=Side(
        style="thin",
        color="BFBFBF"
    ),
    top=Side(
        style="thin",
        color="BFBFBF"
    ),
    bottom=Side(
        style="thin",
        color="BFBFBF"
    )
)
for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.fill = (
            header_fill
        )
        cell.font = (
            header_font
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
    for row in ws.iter_rows():
        for cell in row:
            cell.border = (
                thin_border
            )
            cell.alignment = Alignment(
                vertical="center"
            )
    for column_cells in ws.columns:
        column_letter = (
            column_cells[0]
            .column_letter
        )
        max_length = 0
        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(
                        str(
                            cell.value
                        )
                    )
                )
        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            35
        )
wb.save(
    OUTPUT_EXCEL
)
print(
    "\nExcel file saved:"
)
print(
    OUTPUT_EXCEL
)

# ============================================================

# 31. FIGURE 3.6
# SILHOUETTE SCORE BY NUMBER OF CLUSTERS

# ============================================================

fig, ax = plt.subplots(
    figsize=(8, 5)
)
ax.plot(
    evaluation_df[
        "Number of Clusters (k)"
    ],
    evaluation_df[
        "Silhouette Score"
    ],
    marker="o",
    linewidth=1.8,
    markersize=7
)

# Highlight the optimal k
ax.scatter(
    OPTIMAL_K,
    OPTIMAL_SILHOUETTE,
    s=100,
    zorder=5,
    label=(
        f"Optimal k={OPTIMAL_K} "
        f"({OPTIMAL_SILHOUETTE:.3f})"
    )
)
ax.axvline(
    OPTIMAL_K,
    linestyle="--",
    linewidth=1.0,
    alpha=0.6
)
ax.set_xticks(
    range(
        K_MIN,
        K_MAX + 1
    )
)
ax.set_xlabel(
    "Number of Clusters (k)"
)
ax.set_ylabel(
    "Silhouette Score"
)
ax.set_title(
    "Silhouette Score by Number of Clusters"
)
ax.grid(
    alpha=0.25,
    linestyle="--"
)
ax.legend()
plt.tight_layout()
plt.savefig(
    OUTPUT_G36,
    dpi=300,
    bbox_inches="tight",
    facecolor="white"
)
plt.show()
print(
    "\nFigure 3.6 saved:"
)
print(
    OUTPUT_G36
)

# ============================================================

# 32. TERMINAL OUTPUT

# ============================================================

print("\n" + "=" * 70)
print(
    "FINAL RESULTS — SECTION 3.1.2.2"
)
print("=" * 70)
print(
    "\nTABLE 3.4"
)
print(
    table_34
    .round(4)
    .to_string(
        index=False
    )
)
print(
    "\nOptimal number of clusters:",
    OPTIMAL_K
)
print(
    "Optimal Silhouette Score:",
    round(
        OPTIMAL_SILHOUETTE,
        4
    )
)
print(
    "\nCluster sizes:"
)
print(
    cluster_size
    .round(2)
    .to_string(
        index=False
    )
)
print(
    "\nCluster medoids:"
)
print(
    medoid_df
    .to_string(
        index=False
    )
)
print("=" * 70)
