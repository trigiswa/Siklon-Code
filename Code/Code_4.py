# ============================================================
# 0. IMPORT LIBRARIES
# ============================================================
import os
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)
# ============================================================
# 1. FILE SETTINGS
# ============================================================
FILE_PATH = "Siklon_Tropis_ERA5_Extracted.xlsx"
SHEET_NAME = "Sheet1"
OUTPUT_DIR = "output_3_1_3"
os.makedirs(
    OUTPUT_DIR,
    exist_ok=True
)
OUTPUT_EXCEL = os.path.join(
    OUTPUT_DIR,
    "Genesis_Meteorological_Analysis_3_1_3_1.xlsx"
)
OUTPUT_G39 = os.path.join(
    OUTPUT_DIR,
    "Figure_3.9_Genesis_Meteorological_Distributions.png"
)
# ============================================================
# 2. STUDY DOMAIN
# ============================================================
LON_MIN = 95
LON_MAX = 145
LAT_MIN = -15
LAT_MAX = 25
EQ_MIN = -5
EQ_MAX = 5
YEAR_START = 2001
YEAR_END = 2025
# ============================================================
# 3. ZONE ORDER AND COLORS
# ============================================================
ZONE_ORDER = [
    "North",
    "Equatorial",
    "South"
]
ZONE_COLORS = {
    "North": "#1f77b4",
    "Equatorial": "#ff8c00",
    "South": "#2ca02c"
}
# ============================================================
# 4. VISUAL SETTINGS
# ============================================================
plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 10,
    "axes.titlesize": 11,
    "axes.labelsize": 10,
    "figure.dpi": 150,
    "savefig.dpi": 300
})
# ============================================================
# 5. READ INPUT DATA
# ============================================================
df = pd.read_excel(
    FILE_PATH,
    sheet_name=SHEET_NAME
)
print("=" * 80)
print("INITIAL DATA")
print("=" * 80)
print(
    "Number of rows :",
    len(df)
)
print(
    "Number of SIDs :",
    df["SID"].nunique()
)
# ============================================================
# 6. DATA CLEANING
# ============================================================
df["ISO_TIME"] = pd.to_datetime(
    df["ISO_TIME"],
    errors="coerce"
)
numeric_columns = [
    "LAT",
    "LON",
    "SEASON",
    # Zoom Earth
    "WIND_km/h",
    "PRESSURE_hPa",
    # ERA5
    "era5_sst",
    "era5_tp",
    "era5_w",
    "era5_vo"
]
for col in numeric_columns:
    if col in df.columns:
        df[col] = pd.to_numeric(
            df[col],
            errors="coerce"
        )
# Remove rows with missing essential fields
df = df.dropna(
    subset=[
        "SID",
        "ISO_TIME",
        "LAT",
        "LON"
    ]
).copy()
# Validate latitude
df = df[
    df["LAT"].between(
        -90,
        90
    )
].copy()
# Validate longitude
df = df[
    df["LON"].between(
        -180,
        360
    )
].copy()
# Remove duplicate observations
df = df.drop_duplicates(
    subset=[
        "SID",
        "ISO_TIME"
    ],
    keep="first"
)
# Sort by SID and timestamp
df = df.sort_values(
    [
        "SID",
        "ISO_TIME"
    ]
).reset_index(
    drop=True
)
# ============================================================
# 7. IDENTIFY GENESIS LOCATIONS
# ============================================================
#
# Genesis is operationally defined as
# the first observation for each SID after sorting by time.
# ============================================================
genesis_all = (
    df.drop_duplicates(
        subset="SID",
        keep="first"
    )
    .copy()
)
# ============================================================
# 8. FILTER BY STUDY PERIOD
# ============================================================
genesis_all = genesis_all[
    genesis_all["SEASON"].between(
        YEAR_START,
        YEAR_END
    )
].copy()
# ============================================================
# 9. FILTER STUDY DOMAIN
# ============================================================
genesis = genesis_all[
    genesis_all["LON"].between(
        LON_MIN,
        LON_MAX,
        inclusive="both"
    )
    &
    genesis_all["LAT"].between(
        LAT_MIN,
        LAT_MAX,
        inclusive="both"
    )
].copy()
# ============================================================
# 10. CLASSIFY GENESIS ZONES
# ============================================================
def classify_zone(lat):
    if EQ_MIN <= lat <= EQ_MAX:
        return "Equatorial"
    elif lat > EQ_MAX:
        return "North"
    else:
        return "South"
genesis["Zone"] = (
    genesis["LAT"]
    .apply(
        classify_zone
    )
)
# ============================================================
# 11. POPULATION QUALITY CONTROL
# ============================================================
print("\n" + "=" * 80)
print("GENESIS EVENTS WITHIN THE STUDY DOMAIN")
print("=" * 80)
print(
    "Number of events:",
    len(genesis)
)
print("\nEvent counts by zone:")
print(
    genesis["Zone"]
    .value_counts()
    .reindex(
        ZONE_ORDER,
        fill_value=0
    )
)
# Expected totals for the current study dataset:
#
# Total      = 392
# North      = 354
# Equatorial = 2
# South      = 36
if len(genesis) != 392:
    print(
        "\nWARNING:"
        "\nThe number of events is not 392."
        "\nPlease verify the input file."
    )
# ============================================================
# 12. ERA5 PARAMETER CONVERSIONS
# ============================================================
#
# SST:
# Kelvin -> Celsius
#
# Total precipitation:
# meters -> millimeters
#
# w:
# kept in Pa/s
#
# vo:
# kept in s^-1,
# then scaled by 10^5 for presentation
# ============================================================
# ------------------------------------------------------------
# SST
# ------------------------------------------------------------
genesis["SST_C"] = (
    genesis["era5_sst"]
    -
    273.15
)
# ------------------------------------------------------------
# Total Precipitation
# ------------------------------------------------------------
genesis["TP_mm"] = (
    genesis["era5_tp"]
    *
    1000.0
)
# ------------------------------------------------------------
# Vertical Velocity
# ------------------------------------------------------------
genesis["Omega_Pa_s"] = (
    genesis["era5_w"]
)
# ------------------------------------------------------------
# Relative Vorticity
# ------------------------------------------------------------
genesis["Vorticity_s1"] = (
    genesis["era5_vo"]
)
genesis["Vorticity_1e5"] = (
    genesis["Vorticity_s1"]
    *
    1e5
)
# ============================================================
# 13. ZOOM EARTH PARAMETERS
# ============================================================
#
# Wind and pressure are NOT converted because
# the dataset already uses:
#
# WIND_km/h    -> km/h
# PRESSURE_hPa -> hPa
# ============================================================
genesis["Wind_kmh"] = (
    genesis["WIND_km/h"]
)
genesis["Pressure_hPa"] = (
    genesis["PRESSURE_hPa"]
)
# ============================================================
# 14. CORIOLIS PARAMETER
# ============================================================
#
# f = 2 * Omega * sin(phi)
#
# Omega = Earth's angular rotation rate
# phi   = latitude
# ============================================================
OMEGA_EARTH = 7.2921159e-5
genesis["Coriolis_s1"] = (
    2
    *
    OMEGA_EARTH
    *
    np.sin(
        np.radians(
            genesis["LAT"]
        )
    )
)
# Scaled for table presentation
genesis["Coriolis_1e5"] = (
    genesis["Coriolis_s1"]
    *
    1e5
)
# ============================================================
# 15. DEFINE ANALYSIS PARAMETERS
# ============================================================
PARAMETERS = {
    "SST_C":
        "SST (°C)",
    "TP_mm":
        "Total Precipitation (mm)",
    "Omega_Pa_s":
        "Vertical Velocity, ω (Pa/s)",
    "Vorticity_1e5":
        "Relative Vorticity (×10⁻⁵ s⁻¹)",
    "Wind_kmh":
        "Wind Speed (km/h)",
    "Pressure_hPa":
        "Pressure (hPa)",
    "Coriolis_1e5":
        "Coriolis Parameter (×10⁻⁵ s⁻¹)"
}
# ============================================================
# 16. PARAMETER QUALITY CONTROL
# ============================================================
print("\n" + "=" * 80)
print("PARAMETER QUALITY CONTROL")
print("=" * 80)
for variable, label in PARAMETERS.items():
    n_valid = (
        genesis[variable]
        .notna()
        .sum()
    )
    n_missing = (
        genesis[variable]
        .isna()
        .sum()
    )
    print(
        f"{label:<42}"
        f"Valid = {n_valid:<4} "
        f"Missing = {n_missing}"
    )
# ============================================================
# 17. DATA COMPLETENESS QC BY ZONE
# ============================================================
completeness_records = []
for zone in ZONE_ORDER:
    subset = genesis[
        genesis["Zone"]
        ==
        zone
    ]
    for variable, label in PARAMETERS.items():
        n_total = len(
            subset
        )
        n_valid = (
            subset[variable]
            .notna()
            .sum()
        )
        if n_total > 0:
            completeness = (
                n_valid
                /
                n_total
                *
                100
            )
        else:
            completeness = np.nan
        completeness_records.append({
            "Zone":
                zone,
            "Parameter":
                label,
            "Number of Events":
                n_total,
            "N Valid":
                n_valid,
            "N Missing":
                n_total - n_valid,
            "Completeness (%)":
                completeness
        })
completeness_df = pd.DataFrame(
    completeness_records
)
# ============================================================
# 18. FULL DESCRIPTIVE STATISTICS
# ============================================================
statistics_records = []
for zone in ZONE_ORDER:
    subset = genesis[
        genesis["Zone"]
        ==
        zone
    ]
    for variable, label in PARAMETERS.items():
        values = (
            subset[variable]
            .dropna()
        )
        if len(values) > 0:
            mean_value = (
                values.mean()
            )
            median_value = (
                values.median()
            )
            minimum_value = (
                values.min()
            )
            maximum_value = (
                values.max()
            )
            if len(values) > 1:
                std_value = (
                    values.std()
                )
            else:
                std_value = np.nan
        else:
            mean_value = np.nan
            median_value = np.nan
            minimum_value = np.nan
            maximum_value = np.nan
            std_value = np.nan
        statistics_records.append({
            "Zone":
                zone,
            "Parameter":
                label,
            "N Valid":
                len(values),
            "Mean":
                mean_value,
            "Median":
                median_value,
            "Minimum":
                minimum_value,
            "Maximum":
                maximum_value,
            "Standard Deviation":
                std_value
        })
statistics_df = pd.DataFrame(
    statistics_records
)
# ============================================================
# 19. TABLE 3.7
# MEAN METEOROLOGICAL PARAMETERS AT GENESIS
# ============================================================
table_records = []
for zone in ZONE_ORDER:
    subset = genesis[
        genesis["Zone"]
        ==
        zone
    ]
    table_records.append({
        "Zone":
            zone,
        "Number of Events":
            len(subset),
        "SST (°C)":
            subset["SST_C"].mean(),
        "TP (mm)":
            subset["TP_mm"].mean(),
        "ω (Pa/s)":
            subset["Omega_Pa_s"].mean(),
        "Relative Vorticity (×10⁻⁵ s⁻¹)":
            subset["Vorticity_1e5"].mean(),
        "Wind Speed (km/h)":
            subset["Wind_kmh"].mean(),
        "Pressure (hPa)":
            subset["Pressure_hPa"].mean(),
        "Coriolis (×10⁻⁵ s⁻¹)":
            subset["Coriolis_1e5"].mean()
    })
table_37 = pd.DataFrame(
    table_records
)
# ============================================================
# 20. TWO EQUATORIAL GENESIS CASES
# ============================================================
#
# Saved as supporting data at this stage.
#
# Detailed analysis is carried out in Section 3.1.3.4.
# ============================================================
equatorial_cases = genesis[
    genesis["Zone"]
    ==
    "Equatorial"
][
    [
        "SID",
        "NAME",
        "SEASON",
        "ISO_TIME",
        "LAT",
        "LON",
        "SST_C",
        "TP_mm",
        "Omega_Pa_s",
        "Vorticity_1e5",
        "Wind_kmh",
        "Pressure_hPa",
        "Coriolis_1e5"
    ]
].copy()
# ============================================================
# 21. PRINT EQUATORIAL CASE DATA
# ============================================================
print("\n" + "=" * 80)
print("KEJADIAN GENESIS ZONA EKUATORIAL")
print("=" * 80)
print(
    equatorial_cases[
        [
            "NAME",
            "SEASON",
            "LAT",
            "LON",
            "Wind_kmh",
            "Pressure_hPa"
        ]
    ]
    .to_string(
        index=False
    )
)
# ============================================================
# 22. SAVE RESULTS TO EXCEL
# ============================================================
with pd.ExcelWriter(
    OUTPUT_EXCEL,
    engine="openpyxl"
) as writer:
    # --------------------------------------------------------
    # Main thesis table
    # --------------------------------------------------------
    table_37.round(
        3
    ).to_excel(
        writer,
        sheet_name="Table 3.7",
        index=False
    )
    # --------------------------------------------------------
    # Full statistics
    # --------------------------------------------------------
    statistics_df.round(
        4
    ).to_excel(
        writer,
        sheet_name="Descriptive Statistics",
        index=False
    )
    # --------------------------------------------------------
    # Completeness
    # --------------------------------------------------------
    completeness_df.round(
        2
    ).to_excel(
        writer,
        sheet_name="Data Completeness QC",
        index=False
    )
    # --------------------------------------------------------
    # Two rare equatorial events
    # --------------------------------------------------------
    equatorial_cases.round(
        4
    ).to_excel(
        writer,
        sheet_name="Equatorial Cases",
        index=False
    )
    # --------------------------------------------------------
    # All genesis records
    # --------------------------------------------------------
    genesis.to_excel(
        writer,
        sheet_name="Genesis Data",
        index=False
    )
# ============================================================
# 23. FORMAT EXCEL WORKBOOK
# ============================================================
wb = load_workbook(
    OUTPUT_EXCEL
)
header_fill = PatternFill(
    "solid",
    fgColor="D9EAF7"
)
header_font = Font(
    bold=True
)
thin_border = Border(
    left=Side(
        style="thin",
        color="BFBFBF"
    ),
    right=Side(
        style="thin",
        color="BFBFBF"
    ),
    top=Side(
        style="thin",
        color="BFBFBF"
    ),
    bottom=Side(
        style="thin",
        color="BFBFBF"
    )
)
for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    # Header
    for cell in ws[1]:
        cell.fill = (
            header_fill
        )
        cell.font = (
            header_font
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
    # Body cells
    for row in ws.iter_rows():
        for cell in row:
            cell.border = (
                thin_border
            )
            cell.alignment = Alignment(
                vertical="center"
            )
    # Column widths
    for column_cells in ws.columns:
        column_letter = (
            column_cells[0]
            .column_letter
        )
        max_length = 0
        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(
                        str(
                            cell.value
                        )
                    )
                )
        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            38
        )
wb.save(
    OUTPUT_EXCEL
)
print(
    "\nExcel file successfully saved:"
)
print(
    OUTPUT_EXCEL
)
# ============================================================
# ============================================================
#
#                    FIGURE 3.9
#
# DISTRIBUTION OF METEOROLOGICAL PARAMETERS
# AT GENESIS BY ZONE
#
# ============================================================
# ============================================================
# ============================================================
# 24. PLOT PARAMETERS
# ============================================================
PLOT_PARAMETERS = [
    (
        "SST_C",
        "Sea Surface Temperature",
        "°C"
    ),
    (
        "TP_mm",
        "Total Precipitation",
        "mm"
    ),
    (
        "Omega_Pa_s",
        "Vertical Velocity (ω)",
        "Pa/s"
    ),
    (
        "Vorticity_1e5",
        "Relative Vorticity",
        "×10⁻⁵ s⁻¹"
    ),
    (
        "Wind_kmh",
        "Wind Speed",
        "km/h"
    ),
    (
        "Pressure_hPa",
        "Pressure",
        "hPa"
    ),
    (
        "Coriolis_1e5",
        "Coriolis Parameter",
        "×10⁻⁵ s⁻¹"
    )
]
# ============================================================
# 25. CREATE FIGURE
# ============================================================
fig, axes = plt.subplots(
    2,
    4,
    figsize=(
        18,
        10
    )
)
axes = axes.flatten()
# Random generator used only for point jitter
rng = np.random.default_rng(
    42
)
# ============================================================
# 26. PLOT EACH PARAMETER
# ============================================================
for ax, (
    variable,
    title,
    unit
) in zip(
    axes,
    PLOT_PARAMETERS
):
    # --------------------------------------------------------
    # Data by zone
    # --------------------------------------------------------
    data_by_zone = [
        genesis.loc[
            genesis["Zone"]
            ==
            zone,
            variable
        ]
        .dropna()
        .to_numpy()
        for zone in ZONE_ORDER
    ]
    # --------------------------------------------------------
    # BOXPLOT
    # --------------------------------------------------------
    box = ax.boxplot(
        data_by_zone,
        labels=ZONE_ORDER,
        patch_artist=True,
        widths=0.55,
        showfliers=False,
        medianprops=dict(
            color="black",
            linewidth=1.5
        ),
        whiskerprops=dict(
            linewidth=1.0
        ),
        capprops=dict(
            linewidth=1.0
        )
    )
    # --------------------------------------------------------
    # Box colors
    # --------------------------------------------------------
    for patch, zone in zip(
        box["boxes"],
        ZONE_ORDER
    ):
        patch.set_facecolor(
            ZONE_COLORS[
                zone
            ]
        )
        patch.set_alpha(
            0.42
        )
    # --------------------------------------------------------
    # INDIVIDUAL OBSERVATIONS
    # --------------------------------------------------------
    #
    # Show the two equatorial genesis cases
    # as individual observations.
    # --------------------------------------------------------
    for x_position, zone in enumerate(
        ZONE_ORDER,
        start=1
    ):
        values = (
            genesis.loc[
                genesis["Zone"]
                ==
                zone,
                variable
            ]
            .dropna()
            .to_numpy()
        )
        jitter = rng.normal(
            loc=0,
            scale=0.045,
            size=len(values)
        )
        # Make larger samples more transparent
        if len(values) > 50:
            point_alpha = 0.22
        else:
            point_alpha = 0.70
        ax.scatter(
            np.full(
                len(values),
                x_position
            )
            +
            jitter,
            values,
            s=15,
            color=ZONE_COLORS[
                zone
            ],
            alpha=point_alpha,
            edgecolors="none",
            zorder=3
        )
    # --------------------------------------------------------
    # Axis
    # --------------------------------------------------------
    ax.set_title(
        title,
        fontweight="bold"
    )
    ax.set_ylabel(
        unit
    )
    ax.grid(
        axis="y",
        linestyle="--",
        linewidth=0.5,
        alpha=0.25
    )
# ============================================================
# 27. FINAL PANEL = INFORMATION
# ============================================================
axes[
    7
].axis(
    "off"
)
info_text = (
    f"Number of events: {len(genesis)}\n"
    f"North        = {(genesis['Zone'] == 'North').sum()}\n"
    f"Equatorial   = {(genesis['Zone'] == 'Equatorial').sum()}\n"
    f"South        = {(genesis['Zone'] == 'South').sum()}\n\n"
    "Equatorial cases:\n"
    "VAMEI (2001)\n"
    "SENYAR (2025)\n\n"
    "Data sources:\n"
    "SST, TP, ω, ζ = ERA5\n"
    "Wind, Pressure = Zoom Earth\n"
    "Coriolis = calculated\n\n"
    "Notes:\n"
    "Negative ω = upward motion\n"
    "Dots = individual events\n\n"
    "The equatorial zone has only\n"
    "two events; therefore, results\n"
    "are interpreted descriptively."
)
axes[
    7
].text(
    0.04,
    0.96,
    info_text,
    transform=axes[
        7
    ].transAxes,
    va="top",
    fontsize=10.5,
    linespacing=1.35,
    bbox=dict(
        facecolor="white",
        edgecolor="gray",
        boxstyle="round,pad=0.6",
        alpha=0.95
    )
)
# ============================================================
# 28. MAIN FIGURE TITLE
# ============================================================
fig.suptitle(
    "Meteorological Parameters at Tropical Cyclone Genesis by Zone (2001–2025)",
    fontsize=17,
    fontweight="bold",
    y=0.98
)
# ============================================================
# 29. LAYOUT
# ============================================================
plt.tight_layout(
    rect=[
        0,
        0,
        1,
        0.95
    ]
)
# ============================================================
# 30. SAVE FIGURE
# ============================================================
plt.savefig(
    OUTPUT_G39,
    dpi=300,
    bbox_inches="tight",
    facecolor="white"
)
plt.show()
# ============================================================
# 31. PRINT FINAL RESULTS
# ============================================================
print("\n" + "=" * 110)
print(
    "TABLE 3.7"
)
print(
    "MEAN METEOROLOGICAL PARAMETERS "
    "AT GENESIS BY ZONE"
)
print("=" * 110)
print(
    table_37
    .round(3)
    .to_string(
        index=False
    )
)
print("\n" + "=" * 110)
print(
    "DATA COMPLETENESS"
)
print("=" * 110)
print(
    completeness_df
    .round(2)
    .to_string(
        index=False
    )
)
print(
    "\nAll outputs saved in:"
)
print(
    OUTPUT_DIR
)
print("=" * 110)
